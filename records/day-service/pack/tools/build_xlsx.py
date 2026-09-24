import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1]
F = '游ゴシック'
def font(b=False, s=10, c='000000'): return Font(name=F, bold=b, size=s, color=c)
thin = Side(style='thin', color='808080'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='D9D9D9'); INPUT = PatternFill('solid', fgColor='FFF9DB')
RED = PatternFill('solid', fgColor='F8CBAD'); ORG = PatternFill('solid', fgColor='FFE699'); GRN = PatternFill('solid', fgColor='E2EFDA')
C = Alignment(horizontal='center', vertical='center', wrap_text=True); L = Alignment(horizontal='left', vertical='center', wrap_text=True)
DATE = 'yyyy/m/d'
N = 50  # 職員の行数
METHODS = '"集合,オンライン,動画視聴,後日フォロー,その他"'
KUBUN = '"定期研修,新規採用時研修,その他"'

wb = Workbook()

def setup(ws, landscape=True):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.6; ws.page_margins.top = ws.page_margins.bottom = 0.7
    ws.sheet_view.showGridLines = False

def cell(ws, ref, v, f=None, fill=None, al=L, fmt=None, border=True):
    c = ws[ref]; c.value = v; c.font = f or font(); c.alignment = al if border else Alignment(vertical='center', wrap_text=False)
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if border: c.border = BD
    return c

# ---------- 受講管理表 ----------
def kanri(ws, ex=None):
    setup(ws)
    title = '【記入例】' if ex else ''
    cell(ws, 'A1', f'{title}感染症の予防及びまん延の防止のための研修　職員別 受講管理表（通所介護・デイサービス）', font(True, 13), border=False)
    cell(ws, 'A2', '黄色の欄に入力してください。白・緑の欄は自動で計算されます。' + ('（職員名・日付はすべて架空です）' if ex else ''), font(s=9), border=False)
    cell(ws, 'A3', '事業所名', font(True), HEAD); ws.merge_cells('B3:D3'); cell(ws, 'B3', ex['jigyosho'] if ex else None, fill=INPUT)
    for col in 'CD': ws[f'{col}3'].border = BD
    cell(ws, 'A4', '年度開始日', font(True), HEAD); cell(ws, 'B4', ex['start'] if ex else None, fill=INPUT, fmt=DATE)
    cell(ws, 'C4', '年度終了日', font(True), HEAD); cell(ws, 'D4', '=IF(B4="","",EDATE(B4,12)-1)', fmt=DATE)
    # 集計
    r0, r1 = 8, 8 + N - 1
    labels = [('在籍職員数', f'=COUNTA(C{r0}:C{r1})', '0"名"'),
              ('定期研修 受講済', f'=COUNTIF(K{r0}:K{r1},"受講済")', '0"名"'),
              ('定期研修 未受講', f'=COUNTIF(K{r0}:K{r1},"未受講")', '0"名"'),
              ('定期研修 受講率', f'=IF(F3=0,"",G3/F3)', '0%'),
              ('新規採用時研修 要確認', f'=COUNTIF(L{r0}:L{r1},"要確認")', '0"名"')]
    heads = 'FGHIJ'
    for i, (lab, fml, fmt) in enumerate(labels):
        cell(ws, f'{heads[i]}2', lab, font(True, 9), HEAD, C); ws.row_dimensions[2].height = 30
        cell(ws, f'{heads[i]}3', fml.replace('F3', 'F3').replace('G3', 'G3'), font(True, 12), GRN, C, fmt)
    # 見出し
    hdr = ['No.', '職種', '氏名', '採用日', '新規採用時研修\n受講日', '定期研修①\n受講日', '①の方法', '定期研修②\n受講日', '②の方法',
           '欠席時フォロー\n実施日', '定期研修の状況\n（自動）', '新規採用時研修の状況\n（自動）', '備考']
    widths = [5, 11, 14, 11, 13, 12, 11, 12, 11, 13, 13, 15, 26]
    for i, (h, w) in enumerate(zip(hdr, widths)):
        col = chr(65 + i); ws.column_dimensions[col].width = w
        cell(ws, f'{col}7', h, font(True, 9), HEAD, C)
    ws.row_dimensions[7].height = 32
    rows = ex['rows'] if ex else []
    for k in range(N):
        r = r0 + k
        v = rows[k] if k < len(rows) else [None] * 10
        cell(ws, f'A{r}', k + 1, al=C)
        for j, col in enumerate('BCDEFGHIJ'):
            fmt = DATE if col in 'DEFHJ' else None
            cell(ws, f'{col}{r}', v[j], fill=INPUT, fmt=fmt, al=C if col != 'C' else L)
        cell(ws, f'K{r}', f'=IF(C{r}="","",IF(COUNT(F{r},H{r},J{r})>0,"受講済","未受講"))', al=C)
        cell(ws, f'L{r}', f'=IF(OR(C{r}="",D{r}="",$B$4=""),"",IF(D{r}<$B$4,"対象外",IF(E{r}<>"","受講済","要確認")))', al=C)
        cell(ws, f'M{r}', v[9], fill=INPUT)
    ws.conditional_formatting.add(f'K{r0}:K{r1}', FormulaRule(formula=[f'K{r0}="未受講"'], fill=RED))
    ws.conditional_formatting.add(f'L{r0}:L{r1}', FormulaRule(formula=[f'L{r0}="要確認"'], fill=ORG))
    dv = DataValidation(type='list', formula1=METHODS, allow_blank=True); ws.add_data_validation(dv)
    dv.add(f'G{r0}:G{r1}'); dv.add(f'I{r0}:I{r1}')
    ws.freeze_panes = 'D8'
    ws.print_title_rows = '7:7'
    ws['K7'].comment = Comment('定期研修①・②・欠席時フォローのどれかに日付があれば「受講済」、なければ「未受講」と表示します。', '様式')
    ws['L7'].comment = Comment('採用日が年度開始日以降の職員について、新規採用時研修の受講日がなければ「要確認」と表示します。', '様式')

# ---------- 研修実施一覧 ----------
def ichiran(ws, ex=None, kanri_name='受講管理表'):
    setup(ws)
    title = '【記入例】' if ex else ''
    cell(ws, 'A1', f'{title}感染症の予防及びまん延の防止のための研修　研修実施一覧（年度）', font(True, 13), border=False)
    cell(ws, 'A2', '年度開始日・終了日は「' + kanri_name + '」シートから自動で表示されます。', font(s=9), border=False)
    cell(ws, 'A3', '年度', font(True), HEAD)
    cell(ws, 'B3', f"=IF('{kanri_name}'!B4=\"\",\"\",'{kanri_name}'!B4)", fmt=DATE, al=C)
    cell(ws, 'C3', f"=IF('{kanri_name}'!D4=\"\",\"\",'{kanri_name}'!D4)", fmt=DATE, al=C)
    cell(ws, 'E3', '年度内の定期研修の実施回数', font(True, 9), HEAD, C); ws.merge_cells('E3:F3'); ws['F3'].border = BD
    r0, r1 = 7, 30
    cell(ws, 'G3', f'=IF(B3="","",COUNTIFS(D{r0}:D{r1},"定期研修",B{r0}:B{r1},">="&B3,B{r0}:B{r1},"<="&C3))', font(True, 12), GRN, C, '0"回"')
    cell(ws, 'H3', '※必要な回数は各自治体の通知を確認してください', font(s=9), border=False)
    hdr = ['回', '実施日', '研修名', '研修区分', '実施方法', '講師・進行担当者', '参加者数', '記録作成日', '委員会への報告日', '備考']
    widths = [5, 11, 30, 13, 12, 22, 9, 11, 13, 26]
    for i, (h, w) in enumerate(zip(hdr, widths)):
        col = chr(65 + i); ws.column_dimensions[col].width = w
        cell(ws, f'{col}6', h, font(True, 9), HEAD, C)
    rows = ex['ichiran'] if ex else []
    for k in range(r1 - r0 + 1):
        r = r0 + k; v = rows[k] if k < len(rows) else [None] * 9
        cell(ws, f'A{r}', k + 1, al=C)
        for j, col in enumerate('BCDEFGHIJ'):
            cell(ws, f'{col}{r}', v[j], fill=INPUT, fmt=DATE if col in 'BHI' else ('0"名"' if col == 'G' else None), al=L if col in 'CFJ' else C)
    for ref, f1 in ((f'D{r0}:D{r1}', KUBUN), (f'E{r0}:E{r1}', METHODS)):
        dv = DataValidation(type='list', formula1=f1, allow_blank=True); ws.add_data_validation(dv); dv.add(ref)
    ws.freeze_panes = 'A7'

d = dt.date
EX = {
    'jigyosho': 'サンプル通所介護事業所A（架空）', 'start': d(2026, 4, 1),
    'rows': [
        ['管理者', '職員A', d(2019, 4, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['看護職員', '職員B', d(2020, 10, 1), None, d(2026, 6, 10), '集合', None, None, None, '6月研修の講師'],
        ['介護職員', '職員C', d(2021, 4, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['生活相談員', '職員D', d(2022, 7, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['介護職員', '職員E', d(2023, 4, 1), None, None, None, None, None, d(2026, 6, 15), '6月研修は欠席。録画で後日受講'],
        ['介護職員', '職員F', d(2018, 5, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['介護職員', '職員G', d(2020, 4, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['機能訓練指導員', '職員H', d(2021, 11, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['介護職員', '職員I', d(2022, 4, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['介護職員', '職員J', d(2024, 2, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['看護職員', '職員K', d(2024, 9, 1), None, d(2026, 6, 10), '集合', None, None, None, None],
        ['介護職員', '職員L', d(2025, 6, 1), None, None, None, None, None, d(2026, 6, 17), '6月研修は欠席。録画で後日受講'],
        ['介護職員', '職員M', d(2026, 8, 1), d(2026, 8, 3), None, None, None, None, None, '定期研修は12月に受講予定'],
        ['介護職員', '職員N', d(2026, 9, 1), None, None, None, None, None, None, '9月採用。新規採用時研修を調整中'],
    ],
    'ichiran': [
        [d(2026, 6, 10), '感染症の予防及びまん延の防止のための研修（定期）', '定期研修', '集合', '看護職員・職員B', 10, d(2026, 6, 12), d(2026, 6, 24), '欠席者2名は6/15・6/17に録画で受講'],
        [d(2026, 8, 3), '新規採用時の感染対策研修', '新規採用時研修', '集合', '看護職員・職員B', 1, d(2026, 8, 3), d(2026, 9, 24), '職員M'],
    ],
}

ws = wb.active; ws.title = '使い方'; setup(ws, False)
ws.column_dimensions['A'].width = 100
lines = [
    ('感染症の予防及びまん延の防止のための研修　職員別 受講管理表（通所介護・デイサービス）', font(True, 13)),
    ('', None),
    ('■ シートの構成', font(True, 11)),
    ('・受講管理表：職員ごとに、1年間の研修の受講日・方法・欠席時のフォローを記録します', None),
    ('・研修実施一覧：年度内に実施した研修を1行ずつ記録し、定期研修の実施回数を自動で数えます', None),
    ('・記入例_受講管理表／記入例_研修実施一覧：架空の職員で入力した例です', None),
    ('・免責事項：必ずお読みください', None),
    ('', None),
    ('■ 使い方', font(True, 11)),
    ('1. 「受講管理表」の黄色の欄に、事業所名と年度開始日（例：2026/4/1）を入力します', None),
    ('2. 職員の職種・氏名・採用日を入力します（50名分あります）', None),
    ('3. 研修を行ったら、受講日と方法を入力します。方法はプルダウンから選べます', None),
    ('4. 欠席した職員に後日フォローした場合は「欠席時フォロー実施日」に入力します', None),
    ('5. 「定期研修の状況」に「未受講」（赤）、「新規採用時研修の状況」に「要確認」（橙）が出た職員を確認します', None),
    ('6. 研修ごとの記録は、別の記録用紙（Word・PDF）に作成し、この表とあわせて保管してください', None),
    ('', None),
    ('■ 自動表示のしくみ', font(True, 11)),
    ('・定期研修の状況：定期研修①・②・欠席時フォローのどれかに日付があれば「受講済」、なければ「未受講」', None),
    ('・新規採用時研修の状況：採用日が年度開始日以降の職員で、新規採用時研修の受講日がなければ「要確認」。年度開始日より前の採用は「対象外」', None),
    ('・この表示は入力内容から機械的に判定するものです。研修の実施回数や対象者の考え方は、各自治体の通知を確認してください', None),
    ('', None),
    ('■ 実在の個人情報の取り扱い', font(True, 11)),
    ('・実際に使う場合は職員の氏名を入力するため、事業所の個人情報保護のルールに従って管理してください', None),
    ('・見本として共有・配布する場合は、実在する職員の氏名を入力しないでください', None),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f or font(); c.alignment = Alignment(wrap_text=True, vertical='top')

kanri(wb.create_sheet('受講管理表'))
ichiran(wb.create_sheet('研修実施一覧'))
kanri(wb.create_sheet('記入例_受講管理表'), EX)
ichiran(wb.create_sheet('記入例_研修実施一覧'), EX, '記入例_受講管理表')

ws = wb.create_sheet('免責事項'); setup(ws, False); ws.column_dimensions['A'].width = 100
disc = [
    ('免責事項', font(True, 13)),
    ('・本様式は2026年9月時点で作成者が確認できた情報をもとに作成しています。法令・通知は改正されます。ご利用の際は、最新の法令・通知と各自治体の条例・通知を、ご自身で必ず確認してください。', font(True)),
    ('・本様式は記録作成・管理の参考として作成したひな形であり、法令への適合、監査・運営指導への対応、報酬の減算の回避、立入検査の通過を保証するものではありません。', font(True)),
    ('・作成者は行政書士・社会保険労務士等の専門資格者ではなく、本様式は法的助言を目的とするものではありません。', None),
    ('・自動表示（受講済・未受講・要確認・実施回数）は入力内容から機械的に判定するものであり、基準を満たしているかどうかを判定するものではありません。', None),
    ('・記入例の事業所名・職員名・日付はすべて架空のものであり、実在する施設・店舗・個人・自治体とは関係ありません。', None),
    ('・見本として共有・配布する場合は、実在する利用者・職員の個人情報を入力しないでください。', font(True)),
    ('・本様式の利用により生じたいかなる損害についても、作成者は責任を負いかねます。', None),
]
for i, (t, f) in enumerate(disc, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f or font(); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[i].height = 36 if i > 1 else 22

wb.properties.creator = '記録様式ひな形'
wb.properties.title = '通所介護 デイサービス 感染症研修 職員別 受講管理表 Excel 記入例付き'
wb.properties.keywords = '通所介護 デイサービス 感染症 研修 受講管理表 記録 Excel ひな形'
wb.active = 1
wb.save(OUT)
